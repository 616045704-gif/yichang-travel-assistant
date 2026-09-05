#!/usr/bin/env python3
"""Convert a user-supplied location workbook to ignored CloudBase JSON Lines files."""

import argparse
import hashlib
import json
import pathlib
import re
import sys
from collections import Counter
from datetime import datetime, timezone

from openpyxl import load_workbook


EXPECTED_HEADERS = ['类别', '餐类', '名称', '所属景区', '城市/区县', '地址', '经纬度', '电话', '内容', '攻略/建议', '价格参考', '营业时间', '设施/服务', '车位信息', '标签', '信息更新时间', '数据来源']
CATEGORY_MAP = {'景区概览': 'scenic', '子景点': 'scenic', '博物馆/文化馆': 'culture', '露营地': 'camping', '餐馆': 'restaurant'}
REFERENCE_NOTICE = '资料参考，出行前请以官方公告为准。'
REPO_ROOT = pathlib.Path(__file__).resolve().parents[1]
IMPORT_ROOT = (REPO_ROOT / '.local' / 'import').resolve()


def clean(value, limit=4000):
    if value is None:
        return ''
    text = str(value).replace('\x00', '')
    text = re.sub(r'[\x01-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    return re.sub(r'\s+', ' ', text).strip()[:limit]


def parse_coordinate(value):
    parts = [part.strip() for part in clean(value, 80).replace('，', ',').split(',')]
    if len(parts) != 2:
        return None
    try:
        longitude, latitude = float(parts[0]), float(parts[1])
    except ValueError:
        return None
    if not (73 <= longitude <= 135 and 3 <= latitude <= 54):
        return None
    return longitude, latitude


def stable_place_id(row):
    seed = '\0'.join([str(row.get('_row', '')), clean(row.get('名称')), clean(row.get('地址')), clean(row.get('经纬度')), clean(row.get('类别'))])
    return 'place-' + hashlib.sha256(seed.encode('utf-8')).hexdigest()[:16]


def tags(row):
    values = [clean(value, 60) for value in re.split(r'[,，、;；\s]+', clean(row.get('标签'), 400))]
    return list(dict.fromkeys(value for value in values if value))[:12]


def reference_section(row):
    facts = []
    for label, key in [('攻略/建议', '攻略/建议'), ('价格参考', '价格参考'), ('营业时间', '营业时间'), ('设施/服务', '设施/服务'), ('车位信息', '车位信息'), ('联系电话', '电话')]:
        value = clean(row.get(key), 800)
        if value:
            facts.append(f'{label}：{value}')
    return REFERENCE_NOTICE + (' ' + '；'.join(facts) if facts else '')


def convert_workbook_rows(rows, imported_at):
    places, contents, rejected = [], [], []
    category_counts = Counter()
    used_ids = set()
    for index, original in enumerate(rows, start=2):
        row = dict(original)
        row_number = int(row.get('_row', index))
        category = CATEGORY_MAP.get(clean(row.get('类别')))
        name, district, address, content = (clean(row.get(key), limit) for key, limit in [('名称', 120), ('城市/区县', 120), ('地址', 240), ('内容', 3200)])
        coordinate = parse_coordinate(row.get('经纬度'))
        if not category:
            rejected.append({'row': row_number, 'reason': '类别不支持'})
            continue
        if not name or not district or not address or not content:
            rejected.append({'row': row_number, 'reason': '必要字段缺失'})
            continue
        if coordinate is None:
            rejected.append({'row': row_number, 'reason': '经纬度无效'})
            continue
        place_id = stable_place_id(row)
        if place_id in used_ids:
            rejected.append({'row': row_number, 'reason': '地点标识重复'})
            continue
        used_ids.add(place_id)
        longitude, latitude = coordinate
        source = {'title': '向半斗整理收集', 'url': None, 'licenseNote': '用户整理参考资料；价格、营业时间等以官方公告为准。', 'verifiedAt': imported_at}
        places.append({'_id': place_id, 'placeId': place_id, 'name': name, 'aliases': [], 'category': category, 'district': district, 'address': address, 'latitude': latitude, 'longitude': longitude, 'coordinateSystem': 'GCJ-02', 'intro': content[:280], 'tags': tags(row), 'coverFileId': None, 'sources': [source], 'verifiedAt': imported_at, 'updatedAt': imported_at, 'status': 'published', 'sourceLevel': 'user_collected', 'openNotice': REFERENCE_NOTICE})
        contents.append({'_id': place_id, 'placeId': place_id, 'sections': [{'type': 'text', 'text': content}, {'type': 'text', 'text': reference_section(row)[:4000]}], 'visitAdvice': None, 'diningInfo': None, 'updatedAt': imported_at})
        category_counts[category] += 1
    return {'places': places, 'contents': contents, 'report': {'accepted': len(places), 'rejected': rejected, 'categoryCounts': dict(sorted(category_counts.items()))}}


def safe_output_dir(value):
    output = pathlib.Path(value).resolve() if value else IMPORT_ROOT
    try:
        output.relative_to(IMPORT_ROOT)
    except ValueError as error:
        raise ValueError('输出目录必须位于 .local/import 内。') from error
    output.mkdir(parents=True, exist_ok=True)
    return output


def convert_workbook(source_path, output_dir=None, imported_at=None):
    source = pathlib.Path(source_path).resolve(strict=True)
    if source.suffix.lower() != '.xlsx':
        raise ValueError('仅支持 .xlsx 工作簿。')
    output = safe_output_dir(output_dir)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        header = [clean(value, 120) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        if any(name not in header for name in EXPECTED_HEADERS):
            raise ValueError('首个工作表缺少必要列。')
        rows = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            row = {header[index]: values[index] if index < len(values) else None for index in range(len(header))}
            row['_row'] = row_number
            rows.append(row)
    finally:
        workbook.close()
    timestamp = imported_at or datetime.now(timezone.utc).isoformat(timespec='milliseconds').replace('+00:00', 'Z')
    converted = convert_workbook_rows(rows, timestamp)
    for filename, records in [('places.json', converted['places']), ('place_contents.json', converted['contents'])]:
        (output / filename).write_text(''.join(json.dumps(record, ensure_ascii=False, separators=(',', ':')) + '\n' for record in records), encoding='utf-8')
    return converted


def main(argv=None):
    parser = argparse.ArgumentParser(description='将用户地点工作簿转换为 CloudBase JSON Lines。')
    parser.add_argument('--source', required=True)
    parser.add_argument('--out')
    args = parser.parse_args(argv)
    try:
        converted = convert_workbook(args.source, args.out)
        report = converted['report']
        print(json.dumps({'accepted': report['accepted'], 'rejected': len(report['rejected']), 'categoryCounts': report['categoryCounts'], 'rejectedRows': [item['row'] for item in report['rejected']]}, ensure_ascii=False))
        return 0
    except Exception:
        print('转换失败：请检查工作簿格式与必要字段。', file=sys.stderr)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
